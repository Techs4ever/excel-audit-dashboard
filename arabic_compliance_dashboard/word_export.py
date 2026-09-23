"""Official Word export for current-year non-compliance assessment forms."""
from __future__ import annotations

import base64
import io
import re
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = "1F4E79"
BLUE = "2E75B6"
LINE = "9AA5B1"
INK = "1E293B"


def _set_run_font(run, *, size: int = 11, bold: bool = False, color: str | None = None, name: str = "Calibri"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if color:
        h = color.lstrip("#")
        run.font.color.rgb = RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _rtl_paragraph(paragraph, align=WD_ALIGN_PARAGRAPH.RIGHT):
    paragraph.alignment = align
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.15


def _shade(cell, fill: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def _borders(cell, color: str = LINE, sz: str = "8"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _cell_margins(cell, *, top=60, bottom=60, start=80, end=80):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for name, val in (("top", top), ("left", start), ("bottom", bottom), ("right", end)):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def _write_cell(
    cell,
    text: str,
    *,
    fill: str = "FFFFFF",
    bold: bool = False,
    color: str = INK,
    size: int = 11,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
    min_lines: int = 1,
):
    _shade(cell, fill)
    _borders(cell)
    _cell_margins(cell, top=80 if min_lines > 1 else 60, bottom=80 if min_lines > 1 else 60)
    cell.text = ""
    p = cell.paragraphs[0]
    _rtl_paragraph(p, align)
    run = p.add_run(text or "")
    _set_run_font(run, size=size, bold=bold, color=color)
    extra = max(0, min_lines - 1)
    for _ in range(extra):
        p2 = cell.add_paragraph()
        _rtl_paragraph(p2, align)
        _set_run_font(p2.add_run(" "), size=size, color=color)


def _set_table_rtl(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement("w:tblPr")
    bidi = OxmlElement("w:bidiVisual")
    bidi.set(qn("w:val"), "1")
    tblPr.append(bidi)
    grid = OxmlElement("w:tblW")
    grid.set(qn("w:w"), "5000")
    grid.set(qn("w:type"), "pct")
    tblPr.append(grid)


def decode_logo_data_uri(uri: str | None) -> bytes | None:
    raw = str(uri or "").strip()
    if not raw:
        return None
    m = re.match(r"data:image/(?:png|jpe?g|webp|gif);base64,(.+)$", raw, re.I | re.S)
    if not m:
        return None
    try:
        blob = base64.b64decode(m.group(1))
    except Exception:
        return None
    return blob or None


def _no_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _add_document_brand(document: Document, logo_bytes: bytes | None) -> None:
    table = document.add_table(rows=1, cols=2)
    table.autofit = True
    _set_table_rtl(table)
    logo_cell, mark_cell = table.rows[0].cells
    for cell in (logo_cell, mark_cell):
        _no_borders(cell)
        _cell_margins(cell, top=40, bottom=40, start=40, end=40)
    logo_cell.text = ""
    p = logo_cell.paragraphs[0]
    _rtl_paragraph(p, WD_ALIGN_PARAGRAPH.RIGHT)
    if logo_bytes:
        try:
            run = p.add_run()
            run.add_picture(io.BytesIO(logo_bytes), height=Cm(1.7))
        except Exception:
            pass
    mark_cell.text = ""
    mp = mark_cell.paragraphs[0]
    _rtl_paragraph(mp, WD_ALIGN_PARAGRAPH.LEFT)
    run = mp.add_run("إدارة الالتزام")
    _set_run_font(run, size=12, bold=True, color=NAVY)
    rule = document.add_paragraph()
    _rtl_paragraph(rule)
    pPr = rule._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), NAVY)
    pBdr.append(bottom)
    pPr.append(pBdr)


def _pair_row(table, row_idx, left_label, left_value, right_label, right_value):
    """In RTL visual order: right_label | right_value | left_label | left_value."""
    cells = table.rows[row_idx].cells
    _write_cell(cells[0], right_label, bold=True, color=NAVY, size=11)
    _write_cell(cells[1], right_value, size=11)
    _write_cell(cells[2], left_label, bold=True, color=NAVY, size=11)
    _write_cell(cells[3], left_value, size=11)


def _value_span(table, row_idx, value, min_lines: int):
    cells = table.rows[row_idx].cells
    cells[0].merge(cells[3])
    _write_cell(table.rows[row_idx].cells[0], value, size=11, min_lines=min_lines)


def _section_head(table, row_idx, title: str):
    cells = table.rows[row_idx].cells
    cells[0].merge(cells[3])
    _write_cell(
        table.rows[row_idx].cells[0],
        title,
        fill=BLUE,
        bold=True,
        color="FFFFFF",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def _banner(table, row_idx):
    cells = table.rows[row_idx].cells
    cells[0].merge(cells[3])
    _write_cell(
        table.rows[row_idx].cells[0],
        "التقرير التفصيلي لحالات عدم الالتزام والالتزام الجزئي",
        fill=BLUE,
        bold=True,
        color="FFFFFF",
        size=13,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )


def _add_form(document: Document, form: dict[str, str], *, first: bool) -> None:
    if not first:
        document.add_page_break()

    table = document.add_table(rows=16, cols=4)
    table.autofit = True
    _set_table_rtl(table)

    _banner(table, 0)
    _section_head(table, 1, "أولاً: البيانات الأساسية والإطار التنظيمي")
    _pair_row(
        table,
        2,
        "مستوى المخاطر الكامنة",
        form.get("inherent") or "",
        "اسم المشرع",
        form.get("legislator") or "",
    )
    _pair_row(
        table,
        3,
        "حالة الالتزام",
        form.get("compliance_status") or "",
        "اسم النظام",
        form.get("system_name") or "",
    )
    _pair_row(
        table,
        4,
        "رقم المادة",
        form.get("article_no") or "",
        "اسم الهيئة",
        form.get("authority") or "",
    )
    cells = table.rows[5].cells
    cells[1].merge(cells[3])
    _write_cell(cells[0], "اسم اللائحة", bold=True, color=NAVY)
    _write_cell(table.rows[5].cells[1], form.get("regulation") or "")

    _section_head(table, 6, "ثانياً: نص المادة ومتطلبات الالتزام")
    cells = table.rows[7].cells
    cells[0].merge(cells[3])
    _write_cell(table.rows[7].cells[0], "نص المادة / المتطلب النظامي", bold=True, color=NAVY)
    _value_span(table, 8, form.get("legal_text") or "", min_lines=5)
    cells = table.rows[9].cells
    cells[0].merge(cells[3])
    _write_cell(table.rows[9].cells[0], "البنود/المتطلبات غير الملتزم بها", bold=True, color=NAVY)
    _value_span(table, 10, form.get("noncompliant_items") or "", min_lines=3)

    _section_head(table, 11, "ثالثاً: تقييم مخاطر عدم الالتزام والضوابط الرقابية")
    cells = table.rows[12].cells
    cells[0].merge(cells[3])
    _write_cell(table.rows[12].cells[0], "مخاطر عدم الالتزام", bold=True, color=NAVY)
    _value_span(table, 13, form.get("noncompliance_risk") or "", min_lines=4)
    cells = table.rows[14].cells
    cells[0].merge(cells[3])
    _write_cell(table.rows[14].cells[0], "الإجراء / الخطة التصحيحية", bold=True, color=NAVY)
    _value_span(table, 15, form.get("corrective_plan") or "", min_lines=4)

    table2 = document.add_table(rows=2, cols=4)
    table2.autofit = True
    _set_table_rtl(table2)
    _section_head(table2, 0, "رابعاً: خطة التصحيح وإسناد المسؤوليات")
    _pair_row(
        table2,
        1,
        "تاريخ التصحيح المستهدف",
        form.get("target_date") or "",
        "الإدارة المسؤولة عن التنفيذ",
        form.get("department") or "",
    )


def build_assessment_forms_docx(
    forms: list[dict[str, Any]],
    *,
    logo_bytes: bytes | None = None,
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)

    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    if not forms:
        p = document.add_paragraph()
        _rtl_paragraph(p, WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run("لا توجد سجلات مطابقة للتقرير التفصيلي.")
        _set_run_font(run, size=14, bold=True, color=NAVY)
    else:
        for i, form in enumerate(forms):
            _add_form(document, {str(k): str(v or "") for k, v in (form or {}).items()}, first=i == 0)

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_legal_text_docx(
    legal_text: str,
    fields: list[dict[str, Any]] | None = None,
    *,
    logo_bytes: bytes | None = None,
    title: str = "النص النظامي",
) -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.4)
    section.right_margin = Cm(1.4)
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    heading = document.add_paragraph()
    _rtl_paragraph(heading, WD_ALIGN_PARAGRAPH.RIGHT)
    run = heading.add_run(title)
    _set_run_font(run, size=20, bold=True, color=NAVY)

    legislator = ""
    rows: list[tuple[str, str]] = []
    for field in fields or []:
        label = str(field.get("label") or "").strip()
        value = str(field.get("value") or "").strip()
        if not label and not value:
            continue
        if value in {"", "—", "-", "(blank)"}:
            continue
        if label == "المشرع":
            legislator = value
            continue
        rows.append((label, value))

    if legislator:
        sub = document.add_paragraph()
        _rtl_paragraph(sub, WD_ALIGN_PARAGRAPH.RIGHT)
        run = sub.add_run(f"المشرع: {legislator}")
        _set_run_font(run, size=12, bold=True, color=NAVY)

    body = document.add_paragraph()
    _rtl_paragraph(body, WD_ALIGN_PARAGRAPH.RIGHT)
    run = body.add_run(legal_text or "—")
    _set_run_font(run, size=12, color=INK)

    if rows:
        table = document.add_table(rows=1 + len(rows), cols=2)
        table.autofit = True
        _set_table_rtl(table)
        _write_cell(
            table.rows[0].cells[0],
            "الحقل",
            fill=NAVY,
            bold=True,
            color="FFFFFF",
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        _write_cell(
            table.rows[0].cells[1],
            "القيمة",
            fill=NAVY,
            bold=True,
            color="FFFFFF",
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        for i, (label, value) in enumerate(rows, start=1):
            _write_cell(
                table.rows[i].cells[0],
                label,
                fill="EEF2FF",
                bold=True,
                color=NAVY,
                size=10,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )
            _write_cell(
                table.rows[i].cells[1],
                value,
                fill="FFFFFF",
                size=10,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_register_docx(
    records: list[dict[str, Any]],
    *,
    title: str,
    include_years: bool = False,
    logo_bytes: bytes | None = None,
) -> bytes:
    """Landscape Word register used by annual tracking and current-year lists."""
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    heading = document.add_paragraph()
    _rtl_paragraph(heading, WD_ALIGN_PARAGRAPH.RIGHT)
    run = heading.add_run(title)
    _set_run_font(run, size=20, bold=True, color=NAVY)

    sub = document.add_paragraph()
    _rtl_paragraph(sub, WD_ALIGN_PARAGRAPH.RIGHT)
    run = sub.add_run(f"عدد السجلات: {len(records)}")
    _set_run_font(run, size=11, color="64748B")

    headers = [
        ("المشرع", "ECFEFF", NAVY),
        ("رقم المادة", "EEF2FF", "3730A3"),
        ("النص / المتطلب", "DBEAFE", "1E3A8A"),
        ("مستوى المخاطر", NAVY, "FFFFFF"),
        ("الإدارة", "E2E8F0", "334155"),
    ]
    if include_years:
        headers.extend(
            [
                ("السنة السابقة", "FEF3C7", "92400E"),
                ("السنة الحالية", "FEF3C7", "92400E"),
            ]
        )
    cols = len(headers)
    table = document.add_table(rows=1 + max(len(records), 1), cols=cols)
    table.autofit = True
    _set_table_rtl(table)

    for i, (label, fill, color) in enumerate(headers):
        _write_cell(
            table.rows[0].cells[i],
            label,
            fill=fill,
            bold=True,
            color=color,
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    if not records:
        cells = table.rows[1].cells
        cells[0].merge(cells[cols - 1])
        _write_cell(table.rows[1].cells[0], "لا توجد سجلات مطابقة.", align=WD_ALIGN_PARAGRAPH.CENTER)
    else:
        for ridx, rec in enumerate(records, start=1):
            system = str(rec.get("system_name") or "").strip()
            obs = str(rec.get("observation") or "—")
            if system and system != obs:
                obs = f"{obs}\n{system}"
            values = [
                rec.get("legislator") or "—",
                rec.get("article_no") or "—",
                obs,
                rec.get("rating") or "—",
                rec.get("department") or "—",
            ]
            aligns = [
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.RIGHT,
                WD_ALIGN_PARAGRAPH.CENTER,
                WD_ALIGN_PARAGRAPH.RIGHT,
            ]
            if include_years:
                values.extend([rec.get("final_prev") or "—", rec.get("final_curr") or "—"])
                aligns.extend([WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])
            for cidx, (val, align) in enumerate(zip(values, aligns)):
                _write_cell(
                    table.rows[ridx].cells[cidx],
                    str(val),
                    size=10,
                    bold=include_years and cidx == cols - 1,
                    align=align,
                )

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_annual_tracking_docx(
    records: list[dict[str, Any]],
    *,
    logo_bytes: bytes | None = None,
) -> bytes:
    return build_register_docx(
        records,
        title="تقرير تتبع حالة الالتزام السنوي",
        include_years=True,
        logo_bytes=logo_bytes,
    )


def build_assessment_list_docx(
    records: list[dict[str, Any]],
    *,
    logo_bytes: bytes | None = None,
) -> bytes:
    return build_register_docx(
        records,
        title="نتائج تقييم خلال السنة الحالية",
        include_years=False,
        logo_bytes=logo_bytes,
    )


def _hex_color(value: str, fallback: str = "FFFFFF") -> str:
    h = str(value or "").strip().lstrip("#")
    return h if len(h) == 6 else fallback


def build_aging_matrix_docx(
    payload: dict[str, Any],
    *,
    expand_over_year: bool = False,
    logo_bytes: bytes | None = None,
) -> bytes:
    """Landscape Word matrix for ملخص التقادم."""
    data = payload or {}
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    title = document.add_paragraph()
    _rtl_paragraph(title, WD_ALIGN_PARAGRAPH.RIGHT)
    ref = str(data.get("reference") or "").strip()
    heading = f"ملخص التقادم (حتى {ref})" if ref else "ملخص التقادم"
    run = title.add_run(heading)
    _set_run_font(run, size=20, bold=True, color=NAVY)

    risk_cols = list(data.get("risk_columns") or [])
    time_rows = list(data.get("time_rows") or [])
    over_year_rows = list(data.get("over_year_rows") or [])
    col_totals = data.get("column_totals") or {}
    cols = 2 + len(risk_cols)

    display_rows: list[tuple[dict[str, Any], bool]] = []
    for tr in time_rows:
        display_rows.append((tr, False))
        if expand_over_year and str(tr.get("id") or "") == "ge_1y":
            for child in over_year_rows:
                display_rows.append((child, True))

    table = document.add_table(rows=2 + len(display_rows), cols=max(cols, 2))
    table.autofit = True
    _set_table_rtl(table)

    _write_cell(
        table.rows[0].cells[0],
        "الفترة الزمنية",
        fill="BFDBFE",
        bold=True,
        color="0F172A",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for i, rc in enumerate(risk_cols, start=1):
        _write_cell(
            table.rows[0].cells[i],
            str(rc.get("label") or ""),
            fill=_hex_color(rc.get("color"), "7B8794"),
            bold=True,
            color=_hex_color(rc.get("text_color"), "FFFFFF"),
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    _write_cell(
        table.rows[0].cells[cols - 1],
        "المجموع",
        fill="3B82F6",
        bold=True,
        color="FFFFFF",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    def _count_text(value: Any) -> str:
        try:
            n = int(value or 0)
        except (TypeError, ValueError):
            n = 0
        return str(n) if n else ""

    for ridx, (tr, is_child) in enumerate(display_rows, start=1):
        cells = tr.get("cells") or {}
        time_fill = "F0F9FF" if is_child else "E0F2FE"
        time_color = "1E3A8A" if is_child else "0F172A"
        _write_cell(
            table.rows[ridx].cells[0],
            str(tr.get("label") or ""),
            fill=time_fill,
            bold=not is_child,
            color=time_color,
            size=10,
            align=WD_ALIGN_PARAGRAPH.RIGHT,
        )
        for i, rc in enumerate(risk_cols, start=1):
            rid = str(rc.get("id") or "")
            _write_cell(
                table.rows[ridx].cells[i],
                _count_text(cells.get(rid, 0)),
                fill="FFFFFF",
                bold=True,
                color="1D4ED8",
                size=11,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        _write_cell(
            table.rows[ridx].cells[cols - 1],
            _count_text(tr.get("total")),
            fill="3B82F6",
            bold=True,
            color="FFFFFF",
            size=11,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    foot = table.rows[1 + len(display_rows)]
    _write_cell(
        foot.cells[0],
        "المجموع",
        fill="E0F2FE",
        bold=True,
        color="0F172A",
        size=10,
        align=WD_ALIGN_PARAGRAPH.RIGHT,
    )
    for i, rc in enumerate(risk_cols, start=1):
        rid = str(rc.get("id") or "")
        _write_cell(
            foot.cells[i],
            _count_text(col_totals.get(rid, 0)),
            fill="3B82F6",
            bold=True,
            color="FFFFFF",
            size=11,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
    _write_cell(
        foot.cells[cols - 1],
        _count_text(data.get("grand_total")),
        fill="3B82F6",
        bold=True,
        color="FFFFFF",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_plan_status_docx(
    payload: dict[str, Any],
    *,
    logo_bytes: bytes | None = None,
) -> bytes:
    """Landscape Word report for department corrective-plan status."""
    data = payload or {}
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.1)
    section.right_margin = Cm(1.1)
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    title = document.add_paragraph()
    _rtl_paragraph(title, WD_ALIGN_PARAGRAPH.RIGHT)
    heading = str(
        data.get("title")
        or "حالة خطط المعالجة والإجراءات التصحيحية المتفق عليها مع الإدارة."
    )
    run = title.add_run(heading)
    _set_run_font(run, size=16, bold=True, color=NAVY)
    run.underline = True

    ref = str(data.get("reference") or "").strip()
    if ref:
        meta = document.add_paragraph()
        _rtl_paragraph(meta, WD_ALIGN_PARAGRAPH.RIGHT)
        meta_run = meta.add_run(f"تاريخ المرجع: {ref}")
        _set_run_font(meta_run, size=10, bold=True, color="64748B")

    risk_cols = list(data.get("risk_columns") or [])
    departments = list(data.get("departments") or [])
    totals = data.get("totals") or {}
    group_width = len(risk_cols) + 1
    cols = 2 + (group_width * 2)
    table = document.add_table(rows=3 + max(len(departments), 1), cols=max(cols, 4))
    table.autofit = True
    _set_table_rtl(table)

    def _count_text(value: Any) -> str:
        try:
            n = int(value or 0)
        except (TypeError, ValueError):
            n = 0
        return str(n) if n else ""

    _write_cell(
        table.rows[0].cells[0],
        "الادارة او الجهة المعنية",
        fill="F8FAFC",
        bold=True,
        color="0F172A",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _write_cell(
        table.rows[0].cells[1],
        "عدد النصوص النظامية",
        fill="F8FAFC",
        bold=True,
        color="0F172A",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _write_cell(
        table.rows[0].cells[2],
        "مفتوحة ضمن الجدول الزمني",
        fill="3D7A5A",
        bold=True,
        color="FFFFFF",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _write_cell(
        table.rows[0].cells[2 + group_width],
        "مفتوحة تجاوزت الجدول الزمني",
        fill="8F1D2C",
        bold=True,
        color="FFFFFF",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).merge(table.cell(1, 1))
    table.cell(0, 2).merge(table.cell(0, 1 + group_width))
    table.cell(0, 2 + group_width).merge(table.cell(0, cols - 1))

    for offset in (2, 2 + group_width):
        for i, rc in enumerate(risk_cols):
            _write_cell(
                table.rows[1].cells[offset + i],
                str(rc.get("label") or ""),
                fill=_hex_color(rc.get("color"), "7B8794"),
                bold=True,
                color=_hex_color(rc.get("text_color"), "FFFFFF"),
                size=9,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        _write_cell(
            table.rows[1].cells[offset + len(risk_cols)],
            "الاجمالي",
            fill="5B99C9",
            bold=True,
            color="FFFFFF",
            size=9,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    body_start = 2
    if not departments:
        _write_cell(
            table.rows[body_start].cells[0],
            "لا توجد خطط مفتوحة مطابقة.",
            fill="FFFFFF",
            bold=True,
            color="64748B",
            size=11,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        table.cell(body_start, 0).merge(table.cell(body_start, cols - 1))
        foot_idx = body_start + 1
    else:
        for ridx, dept in enumerate(departments):
            row = table.rows[body_start + ridx]
            fill = "F8FAFC" if ridx % 2 else "FFFFFF"
            _write_cell(
                row.cells[0],
                str(dept.get("label") or "—"),
                fill=fill,
                bold=True,
                color="0F172A",
                size=10,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
            )
            _write_cell(
                row.cells[1],
                _count_text(dept.get("legal_text_count")),
                fill=fill,
                bold=True,
                color="1F4E79",
                size=11,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            for offset, group in ((2, "within"), (2 + group_width, "overdue")):
                cells = dept.get(group) or {}
                for i, rc in enumerate(risk_cols):
                    _write_cell(
                        row.cells[offset + i],
                        _count_text(cells.get(rc.get("id"), 0)),
                        fill=fill,
                        bold=True,
                        color="1E293B",
                        size=11,
                        align=WD_ALIGN_PARAGRAPH.CENTER,
                    )
                _write_cell(
                    row.cells[offset + len(risk_cols)],
                    _count_text(dept.get(f"{group}_total")),
                    fill="E0F2FE",
                    bold=True,
                    color="0F172A",
                    size=11,
                    align=WD_ALIGN_PARAGRAPH.CENTER,
                )
        foot_idx = body_start + len(departments)

    foot = table.rows[foot_idx]
    _write_cell(
        foot.cells[0],
        "الاجمالي",
        fill="1F4E79",
        bold=True,
        color="FFFFFF",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _write_cell(
        foot.cells[1],
        _count_text(totals.get("legal_text_count")),
        fill="1F4E79",
        bold=True,
        color="FFFFFF",
        size=11,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    for offset, group in ((2, "within"), (2 + group_width, "overdue")):
        cells = totals.get(group) or {}
        for i, rc in enumerate(risk_cols):
            _write_cell(
                foot.cells[offset + i],
                _count_text(cells.get(rc.get("id"), 0)),
                fill="1E3A5F",
                bold=True,
                color="FFFFFF",
                size=11,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
        _write_cell(
            foot.cells[offset + len(risk_cols)],
            _count_text(totals.get(f"{group}_total")),
            fill="163A5C",
            bold=True,
            color="FFFFFF",
            size=11,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def build_program_status_docx(
    payload: dict[str, Any],
    *,
    logo_bytes: bytes | None = None,
) -> bytes:
    """Portrait Word report for overall program compliance status by department."""
    data = payload or {}
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.2)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)
    sectPr = section._sectPr
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    sectPr.append(bidi)
    _add_document_brand(document, logo_bytes)

    title = document.add_paragraph()
    _rtl_paragraph(title, WD_ALIGN_PARAGRAPH.RIGHT)
    run = title.add_run(str(data.get("title") or "الحالة العامة لبرنامج الالتزام"))
    _set_run_font(run, size=18, bold=True, color=NAVY)
    run.underline = True

    guide = document.add_paragraph()
    _rtl_paragraph(guide, WD_ALIGN_PARAGRAPH.RIGHT)
    _set_run_font(guide.add_run(str(data.get("guide") or "دليل تصنيف الحالة العامة للبرنامج")), size=11, bold=True, color="475569")

    levels = list(data.get("levels") or [])
    if levels:
        legend = document.add_table(rows=2, cols=max(len(levels), 1))
        legend.autofit = True
        _set_table_rtl(legend)
        for i, level in enumerate(levels):
            fill = _hex_color(level.get("color"), "3D7A5A")
            _write_cell(
                legend.rows[0].cells[i],
                str(level.get("label") or ""),
                fill=fill,
                bold=True,
                color="FFFFFF",
                size=14,
                align=WD_ALIGN_PARAGRAPH.CENTER,
            )
            _write_cell(
                legend.rows[1].cells[i],
                str(level.get("text") or ""),
                fill="F8FAFC",
                bold=False,
                color="334155",
                size=9,
                align=WD_ALIGN_PARAGRAPH.RIGHT,
                min_lines=3,
            )

    overall = data.get("overall") or {}
    bar = document.add_table(rows=1, cols=2)
    bar.autofit = True
    _set_table_rtl(bar)
    _write_cell(
        bar.rows[0].cells[0],
        "الحالة العامة لبرنامج الالتزام:",
        fill="2A6499",
        bold=True,
        color="FFFFFF",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    _write_cell(
        bar.rows[0].cells[1],
        str(overall.get("label") or "—"),
        fill=_hex_color(overall.get("color"), "3D7A5A"),
        bold=True,
        color="FFFFFF",
        size=14,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )

    caption = document.add_paragraph()
    _rtl_paragraph(caption, WD_ALIGN_PARAGRAPH.RIGHT)
    cap_run = caption.add_run("ملخص المتطلبات النظامية التي خضعت لبرنامج الالتزام حسب الادارة او الجهة المعنية")
    _set_run_font(cap_run, size=10, bold=True, color="334155")

    departments = list(data.get("departments") or [])
    totals = data.get("totals") or {}
    table = document.add_table(rows=2 + max(len(departments), 1), cols=5)
    table.autofit = True
    _set_table_rtl(table)
    headers = [
        ("الادارة او الجهة المعنية", "F1F5F9", "0F172A"),
        ("عدد المتطلبات النظامية", "1F4E79", "FFFFFF"),
        ("عدد المتطلبات الملتزم بها", "3D7A5A", "FFFFFF"),
        ("عدد المتطلبات غير الملتزم بها", "C24141", "FFFFFF"),
        ("عدد المتطلبات الملتزم بها جزئيا", "C9A227", "1E293B"),
    ]
    for i, (label, fill, color) in enumerate(headers):
        _write_cell(
            table.rows[0].cells[i],
            label,
            fill=fill,
            bold=True,
            color=color,
            size=9,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    def _count_text(value: Any) -> str:
        try:
            n = int(value or 0)
        except (TypeError, ValueError):
            n = 0
        return str(n) if n else ""

    if not departments:
        _write_cell(
            table.rows[1].cells[0],
            "لا توجد متطلبات مطابقة.",
            fill="FFFFFF",
            bold=True,
            color="64748B",
            size=11,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )
        table.cell(1, 0).merge(table.cell(1, 4))
        foot_idx = 2
    else:
        for ridx, dept in enumerate(departments):
            row = table.rows[1 + ridx]
            fill = "F8FAFC" if ridx % 2 else "FFFFFF"
            values = [
                (str(dept.get("label") or "—"), fill, "0F172A", WD_ALIGN_PARAGRAPH.RIGHT, True),
                (_count_text(dept.get("legal_text_count")), fill, "1F4E79", WD_ALIGN_PARAGRAPH.CENTER, True),
                (_count_text(dept.get("compliant")), fill, "3D7A5A", WD_ALIGN_PARAGRAPH.CENTER, True),
                (_count_text(dept.get("noncompliant")), fill, "C24141", WD_ALIGN_PARAGRAPH.CENTER, True),
                (_count_text(dept.get("partial")), fill, "92400E", WD_ALIGN_PARAGRAPH.CENTER, True),
            ]
            for i, (text, bg, color, align, bold) in enumerate(values):
                _write_cell(row.cells[i], text, fill=bg, bold=bold, color=color, size=10, align=align)
        foot_idx = 1 + len(departments)

    foot = table.rows[foot_idx]
    foot_vals = [
        ("الاجمالي", "1F4E79"),
        (_count_text(totals.get("legal_text_count")), "1F4E79"),
        (_count_text(totals.get("compliant")), "3D7A5A"),
        (_count_text(totals.get("noncompliant")), "C24141"),
        (_count_text(totals.get("partial")), "C9A227"),
    ]
    for i, (text, fill) in enumerate(foot_vals):
        color = "1E293B" if fill == "C9A227" else "FFFFFF"
        _write_cell(
            foot.cells[i],
            text,
            fill=fill,
            bold=True,
            color=color,
            size=10,
            align=WD_ALIGN_PARAGRAPH.CENTER,
        )

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
